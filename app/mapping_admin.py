import json
import re
from pathlib import Path
from fastapi import APIRouter, Depends, Form, Request
from fastapi.responses import HTMLResponse, RedirectResponse
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session
from fastapi.templating import Jinja2Templates
from .database import get_db
from .models import Backlog, OtlgMapping
from .backlogs import apply_mapping, mapping_key

templates=Jinja2Templates(directory=Path(__file__).parent / "templates")
router=APIRouter()
MAPPING_FILE=Path(r"C:\Development\Rückstände\outputs\otlg-mapping\otlg-textmapping-vollstaendig.xlsx")

def apply_all(db):
 mappings={item.text_key:item for item in db.scalars(select(OtlgMapping)).all()}
 for item in db.scalars(select(Backlog)).all():
  data=json.loads(item.data); apply_mapping(data,mappings.get(mapping_key(data.get("Bemerkung OTLG")))); item.data=json.dumps(data,ensure_ascii=False)

def import_workbook(db,path=MAPPING_FILE):
 workbook=load_workbook(path,read_only=True,data_only=True); sheet=workbook["OTLG Mapping"]
 header_row=next(row for row in sheet.iter_rows(values_only=True) if row and row[0]=="Originaltext")
 header_key=lambda value: re.sub(r"[^a-z0-9]+","",mapping_key(value))
 headers={header_key(value):index for index,value in enumerate(header_row) if value}
 required=["originaltext","avisierterliefertermin","regelautomatisiertesmapping","massnahme","kategoriemanuell"]
 if not all(key in headers for key in required): raise ValueError("Die erforderlichen Mapping-Spalten fehlen.")
 existing={item.text_key:item for item in db.scalars(select(OtlgMapping)).all()}; updated=0
 for row in sheet.iter_rows(min_row=sheet.iter_rows().__next__()[0].row if False else 1,values_only=True):
  text=row[headers["originaltext"]] if len(row)>headers["originaltext"] else None
  if not text or text=="Originaltext": continue
  normalized=mapping_key(text); item=existing.get(normalized)
  if not item: item=OtlgMapping(text_key=normalized,original_text=str(text).strip()); db.add(item); existing[normalized]=item
  rule=row[headers["regelautomatisiertesmapping"]] if len(row)>headers["regelautomatisiertesmapping"] else None
  announced=row[headers["avisierterliefertermin"]] if len(row)>headers["avisierterliefertermin"] else None
  item.mapping_rule=str(rule).strip() if rule else None
  item.announced_delivery=None if item.mapping_rule and "anlagedatum" in item.mapping_rule.casefold() else (str(announced).strip() if announced else None)
  item.measure=str(row[headers["massnahme"]]).strip() if len(row)>headers["massnahme"] and row[headers["massnahme"]] else None
  item.category=str(row[headers["kategoriemanuell"]]).strip() if len(row)>headers["kategoriemanuell"] and row[headers["kategoriemanuell"]] else None
  updated+=1
 apply_all(db); db.commit(); return updated

@router.get("/backlogs/mappings",response_class=HTMLResponse)
def page(request:Request,only_incomplete:str="",notice:str|None=None,db:Session=Depends(get_db)):
 rows=list(db.scalars(select(OtlgMapping).order_by(OtlgMapping.original_text)).all())
 if only_incomplete=="1": rows=[row for row in rows if not row.measure or not row.category]
 return templates.TemplateResponse(request,"mappings.html",{"rows":rows,"only_incomplete":only_incomplete=="1","notice":notice})

@router.post("/backlogs/mappings/import")
def import_mapping(db:Session=Depends(get_db)):
 try: count=import_workbook(db); message=f"{count} Mappingzeilen importiert."
 except Exception as error: db.rollback(); message=f"Import fehlgeschlagen: {error}"
 return RedirectResponse("/backlogs/mappings?notice="+message,303)

@router.post("/backlogs/mappings/{mapping_id}")
def save_mapping(mapping_id:int,announced_delivery:str=Form(""),mapping_rule:str=Form(""),measure:str=Form(""),category:str=Form(""),only_incomplete:str=Form(""),db:Session=Depends(get_db)):
 item=db.get(OtlgMapping,mapping_id)
 if not item: return RedirectResponse("/backlogs/mappings",303)
 item.announced_delivery=announced_delivery.strip() or None; item.mapping_rule=mapping_rule.strip() or None; item.measure=measure.strip() or None; item.category=category.strip() or None
 apply_all(db); db.commit(); return RedirectResponse("/backlogs/mappings?only_incomplete="+only_incomplete,303)




